"""
Service module for AI-powered Kit Complet generation.
Handles document extraction, prompt building, LLM calls, and DOCX rendering.
"""
import logging
import os
from pathlib import Path
from typing import Dict, List, Optional

from django.conf import settings
from django.core.files.base import ContentFile
from django.template.loader import render_to_string
from django.utils import timezone

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

try:
    from openai import OpenAI
    HAVE_OPENAI = True
except ImportError:
    HAVE_OPENAI = False

logger = logging.getLogger(__name__)


def load_kit_consignes() -> str:
    """
    Load the content of the Markdown/text file with instructions for the kit.
    Returns the content as a string.
    """
    # Try multiple possible paths
    possible_paths = [
        Path(settings.BASE_DIR) / "templates" / "ai" / "prompts" / "kit_complet_consigne.md",
        Path(settings.BASE_DIR) / "assets" / "Modèle de Consignes — Kit complet de préparation.md",
        Path(__file__).resolve().parents[2] / "templates" / "ai" / "prompts" / "kit_complet_consigne.md",
    ]
    
    for path in possible_paths:
        if path.exists():
            try:
                return path.read_text(encoding="utf-8")
            except Exception as e:
                logger.warning(f"Error reading consignes from {path}: {e}")
                continue
    
    # Fallback: return a basic template
    logger.warning("Consignes file not found, using fallback template")
    return """# Kit complet de préparation à l'audit

## Objectif
Générer un document structuré de préparation à l'audit basé sur les textes réglementaires fournis.

## Structure attendue
1. Introduction générale
2. Questionnaires de préparation (20 questions max par document + 20 générales)
3. Tableaux d'irrégularités (20 max par document + 10 générales)
4. Synthèse et recommandations
5. Plan d'action

## Instructions
- Analyser chaque texte réglementaire transmis
- Identifier les obligations et points critiques
- Proposer des questions pratiques pour la préparation
- Lister les irrégularités possibles avec solutions
- Fournir un plan d'action priorisé
"""


def extract_text_from_file(file_path: str) -> str:
    """
    Extract text from various file formats.
    Supports: .txt, .md, .docx, .pdf (basic)
    
    Args:
        file_path: Path to the file
        
    Returns:
        Extracted text content
    """
    path = Path(file_path)
    if not path.exists():
        logger.warning(f"File not found: {file_path}")
        return ""
    
    ext = path.suffix.lower()
    
    try:
        if ext in (".txt", ".md"):
            return path.read_text(encoding="utf-8", errors="ignore")
        
        elif ext == ".docx":
            try:
                from docx import Document as DocxDocument
                doc = DocxDocument(str(path))
                return "\n".join([para.text for para in doc.paragraphs])
            except Exception as e:
                logger.warning(f"Error extracting text from DOCX {file_path}: {e}")
                return ""
        
        elif ext == ".pdf":
            # Basic PDF extraction - can be enhanced with PyPDF2 or pdfplumber
            try:
                import PyPDF2
                with open(path, "rb") as f:
                    reader = PyPDF2.PdfReader(f)
                    text = "\n".join([page.extract_text() for page in reader.pages])
                    return text
            except ImportError:
                logger.warning("PyPDF2 not installed, PDF extraction skipped")
                return ""
            except Exception as e:
                logger.warning(f"Error extracting text from PDF {file_path}: {e}")
                return ""
        
        else:
            logger.warning(f"Unsupported file type: {ext}")
            return ""
    
    except Exception as e:
        logger.exception(f"Error extracting text from {file_path}: {e}")
        return ""


def build_prompt_for_kit(
    consignes: str,
    inquiry,
    texts: List[Dict[str, str]]
) -> str:
    """
    Build a comprehensive prompt for the LLM to generate the kit.
    
    Args:
        consignes: The instructions/consignes text
        inquiry: ClientInquiry instance
        texts: List of dicts with 'name' and 'content' keys
        
    Returns:
        Complete prompt string
    """
    # Build client context section
    context_lines = [
        "## CONTEXTE CLIENT",
        f"- Nom & Prénom: {inquiry.contact_name or 'Non renseigné'}",
        f"- Email: {inquiry.email}",
        f"- Organisation: {inquiry.organization_name or 'Non renseigné'}",
        f"- Statut juridique: {inquiry.statut_juridique or 'Non renseigné'}",
        f"- Localisation: {inquiry.location or 'Non renseigné'}",
        f"- Secteur: {inquiry.sector or 'Non renseigné'}",
        f"- Budget: {inquiry.budget_range or 'Non renseigné'}",
    ]
    
    if inquiry.mission_text:
        context_lines.append(f"- Missions: {inquiry.mission_text[:500]}")
    
    if inquiry.context_text:
        context_lines.append(f"- Contexte: {inquiry.context_text[:500]}")
    
    if inquiry.funding_sources:
        context_lines.append(f"- Sources de financement: {', '.join(inquiry.funding_sources)}")
    
    if inquiry.audits_types:
        context_lines.append(f"- Types d'audit: {', '.join(inquiry.audits_types)}")
    
    context_section = "\n".join(context_lines)
    
    # Build documents section
    docs_section = "## DOCUMENTS FOURNIS\n\n"
    if texts:
        for i, text_info in enumerate(texts, 1):
            name = text_info.get("name", f"Document {i}")
            content = text_info.get("content", "")
            # Truncate very long content
            if len(content) > 5000:
                content = content[:5000] + "\n\n[... contenu tronqué ...]"
            docs_section += f"### Document {i}: {name}\n\n{content}\n\n---\n\n"
    else:
        docs_section += "Aucun document fourni.\n\n"
    
    # Combine everything
    prompt = f"""{consignes}

{context_section}

{docs_section}

## INSTRUCTIONS FINALES

Génère un document Markdown structuré en français contenant:
1. Une introduction générale adaptée au contexte client
2. Des questionnaires de préparation (20 questions max par document + 20 générales)
3. Des tableaux d'irrégularités (20 max par document + 10 générales)
4. Une synthèse finale avec recommandations
5. Un plan d'action priorisé

Le document doit être professionnel, clair et directement utilisable pour préparer un audit.
Utilise des titres Markdown (##, ###) pour structurer le contenu.
"""
    
    return prompt


def call_llm_kit_complet(prompt: str) -> tuple[str, str, int | None]:
    """
    Call OpenAI (or similar LLM) to generate the kit content.
    
    Args:
        prompt: The complete prompt for generation
        
    Returns:
        Tuple of (generated_markdown, model_name, token_usage)
    """
    if not HAVE_OPENAI:
        raise ImportError("openai package is not installed")
    
    api_key = getattr(settings, "OPENAI_API_KEY", None) or os.getenv("OPENAI_API_KEY")
    if not api_key:
        raise ValueError("OPENAI_API_KEY not configured in settings or environment")
    
    # Build client kwargs
    client_kwargs = {"api_key": api_key}
    
    if hasattr(settings, "OPENAI_ORG") and settings.OPENAI_ORG:
        client_kwargs["organization"] = settings.OPENAI_ORG
    
    if hasattr(settings, "OPENAI_PROJECT") and settings.OPENAI_PROJECT:
        client_kwargs["project"] = settings.OPENAI_PROJECT
    
    if hasattr(settings, "OPENAI_BASE_URL") and settings.OPENAI_BASE_URL:
        client_kwargs["base_url"] = settings.OPENAI_BASE_URL
    
    client = OpenAI(**client_kwargs)
    
    # Determine model
    model = getattr(settings, "OPENAI_CHAT_MODEL", "gpt-4o-mini")
    fallbacks = getattr(
        settings,
        "OPENAI_CHAT_MODEL_FALLBACKS",
        ["gpt-4o", "gpt-4o-mini"]
    )
    models_to_try = [model] + [f for f in fallbacks if f != model]
    
    messages = [
        {
            "role": "system",
            "content": (
                "Tu es un expert en audit et conformité pour les "
                "administrations publiques. Tu génères des documents "
                "structurés et professionnels en français."
            )
        },
        {
            "role": "user",
            "content": prompt
        }
    ]
    
    # Try models with fallback
    response = None
    last_error = None
    used_model = None
    
    for model_name in models_to_try:
        try:
            response = client.chat.completions.create(
                model=model_name,
                messages=messages,
                temperature=0.7,
                max_tokens=8000
            )
            used_model = model_name
            logger.info(f"Successfully used model: {model_name}")
            break
        except Exception as e:
            last_error = e
            logger.warning(f"Error with model {model_name}: {e}")
            if model_name == models_to_try[-1]:
                raise
    
    if response is None:
        raise last_error or ValueError("No model worked")
    
    generated_text = response.choices[0].message.content
    
    # Return model name and token usage info for logging
    token_usage = response.usage.total_tokens if hasattr(response, "usage") else None
    
    return generated_text, used_model, token_usage


def render_markdown_to_docx(
    markdown_content: str,
    template_path: Optional[str] = None,
    inquiry=None
) -> Document:
    """
    Convert Markdown content to a Word Document.
    If a template is provided, loads it and appends content.
    Otherwise, creates a fresh document.
    
    Args:
        markdown_content: The Markdown text to convert
        template_path: Optional path to a Word template
        inquiry: Optional ClientInquiry instance for metadata
        
    Returns:
        python-docx Document object
    """
    # Load template if provided
    if template_path and Path(template_path).exists():
        try:
            doc = Document(template_path)
        except Exception as e:
            logger.warning(f"Error loading template {template_path}: {e}, creating new document")
            doc = Document()
    else:
        doc = Document()
    
    # Add title if document is empty or we want to add a header
    if not doc.paragraphs or len(doc.paragraphs) == 0:
        title = doc.add_heading("Kit complet de préparation à l'audit", 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Add metadata paragraph if inquiry provided
    if inquiry:
        meta_para = doc.add_paragraph()
        meta_para.add_run(f"Préparé pour: {inquiry.contact_name or 'Client'}").bold = True
        if inquiry.organization_name:
            meta_para.add_run(f" ({inquiry.organization_name})")
        meta_para.add_run(f"\nDate: {timezone.now().strftime('%d/%m/%Y')}")
        doc.add_paragraph()  # Spacing
    
    # Parse Markdown and convert to DOCX
    # Basic Markdown parsing (can be enhanced with markdown library)
    lines = markdown_content.split("\n")
    current_para = None
    
    for line in lines:
        line = line.strip()
        
        if not line:
            if current_para:
                current_para = None
            continue
        
        # Headings
        if line.startswith("# "):
            doc.add_heading(line[2:], level=1)
            current_para = None
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=2)
            current_para = None
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=3)
            current_para = None
        elif line.startswith("#### "):
            doc.add_heading(line[5:], level=4)
            current_para = None
        
        # Lists
        elif line.startswith("- ") or line.startswith("* "):
            if not current_para:
                current_para = doc.add_paragraph()
            current_para.add_run(line[2:]).font.size = Pt(11)
            current_para.style = "List Bullet"
        
        elif line.startswith("1. ") or (line[0].isdigit() and ". " in line[:5]):
            if not current_para:
                current_para = doc.add_paragraph()
            current_para.add_run(line.split(". ", 1)[1] if ". " in line else line).font.size = Pt(11)
            current_para.style = "List Number"
        
        # Regular paragraph
        else:
            if current_para:
                current_para.add_run(" " + line)
            else:
                current_para = doc.add_paragraph(line)
                current_para.style.font.size = Pt(11)
    
    return doc


def generate_kit_draft_for_inquiry(inquiry) -> "GeneratedDraft":
    """
    Main orchestration function used by Celery task.
    
    - Loads consignes
    - Extracts text from all attachments
    - Builds prompt
    - Calls LLM
    - Renders Markdown -> DOCX
    - Saves GeneratedDraft and updates inquiry state
    
    Args:
        inquiry: ClientInquiry instance
        
    Returns:
        GeneratedDraft instance
    """
    from store.models import GeneratedDraft, InquiryDocument
    
    logger.info(f"Starting kit generation for inquiry {inquiry.pk}")
    
    try:
        # 1. Load consignes
        consignes = load_kit_consignes()
        
        # 2. Extract texts from attachments
        attachments = InquiryDocument.objects.filter(inquiry=inquiry)
        texts = []
        
        for attachment in attachments:
            try:
                file_path = attachment.file.path
                content = extract_text_from_file(file_path)
                texts.append({
                    "name": attachment.original_name or attachment.file.name,
                    "content": content
                })
            except Exception as e:
                logger.warning(f"Error processing attachment {attachment.id}: {e}")
                # Still add the file name even if extraction fails
                texts.append({
                    "name": attachment.original_name or attachment.file.name,
                    "content": f"[Erreur lors de l'extraction du texte: {e}]"
                })
        
        # 3. Build prompt
        prompt = build_prompt_for_kit(consignes, inquiry, texts)
        
        # 4. Call LLM
        generated_md, model_name, token_usage = call_llm_kit_complet(prompt)
        
        # 5. Render to DOCX
        # Try to find template
        template_path = None
        possible_template_paths = [
            Path(settings.BASE_DIR) / "assets" / "Kit_Complet_Preparation_Couverture_BSG.docx",
            Path(__file__).resolve().parents[2] / "assets" / "Kit_Complet_Preparation_Couverture_BSG.docx",
        ]
        
        for path in possible_template_paths:
            if path.exists():
                template_path = str(path)
                break
        
        doc = render_markdown_to_docx(generated_md, template_path, inquiry)
        
        # 6. Save to file
        from io import BytesIO
        buffer = BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        
        # 7. Create/update GeneratedDraft
        draft, created = GeneratedDraft.objects.get_or_create(
            inquiry=inquiry,
            defaults={
                "model_name": model_name or "",
                "token_usage": token_usage,
                "log": f"Generated at {timezone.now()}\n\nPrompt length: {len(prompt)} chars\nGenerated length: {len(generated_md)} chars"
            }
        )
        
        if not created:
            draft.model_name = model_name or ""
            draft.token_usage = token_usage
            draft.log = f"Regenerated at {timezone.now()}\n\nPrompt length: {len(prompt)} chars\nGenerated length: {len(generated_md)} chars"
        
        # Save DOCX file
        filename = f"kit_{inquiry.pk}_{timezone.now().strftime('%Y%m%d_%H%M%S')}.docx"
        draft.docx.save(filename, ContentFile(buffer.read()), save=True)
        
        # 8. Update inquiry state
        inquiry.processing_state = "DRAFT_DONE"
        inquiry.save(update_fields=["processing_state"])
        
        logger.info(f"Kit generation completed for inquiry {inquiry.pk}")
        
        return draft
    
    except Exception as e:
        logger.exception(f"Error generating kit for inquiry {inquiry.pk}: {e}")
        # Update inquiry state to indicate error
        inquiry.processing_state = "PAID"  # Revert to previous state
        inquiry.save(update_fields=["processing_state"])
        raise

